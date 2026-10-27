"""Text extraction for CV and JD files (PDF, DOCX, TXT) with a quality score.

PyMuPDF is the primary PDF parser; pdfplumber is the fallback when PyMuPDF is missing or its text scores badly.
Pages are separated by a form feed so later stages can cite page numbers.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

from fastapi import HTTPException

PAGE_BREAK = "\f"
HEADINGS = (
    "summary|profile|objective|experience|work experience|professional experience|employment|education|skills|technical skills|"
    "projects|certifications|publications|awards|languages|interests|references"
)
_HEADING_RE = re.compile(rf"^\s*({HEADINGS})\s*:?\s*$", re.I)
_BULLET_RE = re.compile(r"^\s*([-*•▪●◦‣·]|\d+[.)])\s+")


@dataclass
class Parsed:
    text: str
    quality: int
    warnings: list[str] = field(default_factory=list)
    pages: int = 1
    kind: str = "txt"


def detect_kind(data: bytes, name: str) -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(data)):
        return "docx"
    if name.lower().endswith((".pdf", ".docx")):
        raise HTTPException(400, "That file looks damaged or is not a real PDF or DOCX.")
    return "txt"


def _order_blocks(blocks: list[tuple[float, float, float, float, str]], width: float) -> str:
    """Reading order for one page. Finds a vertical gutter that no narrow block crosses and reads left column, then right."""
    blocks = [b for b in blocks if b[4].strip()]
    if not blocks:
        return ""
    wide = [b for b in blocks if (b[2] - b[0]) > 0.62 * width]
    narrow = [b for b in blocks if b not in wide]
    split: float | None = None
    if len(narrow) >= 4:
        free = [
            x
            for x in range(int(0.2 * width), int(0.8 * width), 4)
            if not any(b[0] < x < b[2] for b in narrow)
            and sum(1 for b in narrow if b[2] <= x) >= 2
            and sum(1 for b in narrow if b[0] >= x) >= 2
        ]
        if free:
            # the middle of the longest unbroken run of free x positions is the gutter
            runs, run = [], [free[0]]
            for x in free[1:]:
                if x - run[-1] <= 4:
                    run.append(x)
                else:
                    runs.append(run)
                    run = [x]
            runs.append(run)
            longest = max(runs, key=len)
            split = (longest[0] + longest[-1]) / 2
    if split is None:
        return "\n".join(b[4].strip() for b in sorted(blocks, key=lambda b: (round(b[1]), b[0])))
    parts: list[str] = []
    ordered_wide = sorted(wide, key=lambda b: b[1])
    cuts = [b[1] for b in ordered_wide]
    bounds = [-1.0, *cuts, 1e9]
    for i in range(len(bounds) - 1):
        lo, hi = bounds[i], bounds[i + 1]
        seg = [b for b in narrow if lo <= b[1] < hi]
        for col in ([b for b in seg if b[2] <= split], [b for b in seg if b[2] > split]):
            parts += [b[4].strip() for b in sorted(col, key=lambda b: b[1])]
        if i < len(ordered_wide):
            parts.append(ordered_wide[i][4].strip())
    return "\n".join(parts)


def _pdf_pymupdf(data: bytes) -> tuple[str, int]:
    import pymupdf

    pages: list[str] = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise HTTPException(400, "That PDF is password protected.")
        for page in doc:
            raw = page.get_text("blocks")
            blocks = [(b[0], b[1], b[2], b[3], b[4]) for b in raw if len(b) > 6 and b[6] == 0]
            pages.append(_order_blocks(blocks, page.rect.width))
        return PAGE_BREAK.join(pages), len(pages)


def _pdf_pdfplumber(data: bytes) -> tuple[str, int]:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [(p.extract_text(layout=False) or "") for p in pdf.pages]
    return PAGE_BREAK.join(pages), len(pages)


def _docx(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise HTTPException(400, "That DOCX file could not be read.") from e
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                if c.text.strip() and c.text.strip() not in cells:
                    cells.append(c.text.strip())
            if cells:
                lines.append("  ".join(cells))
    return "\n".join(lines)


def clean(text: str) -> str:
    """Normalise whitespace, rejoin wrapped lines, and put a blank line before each section heading."""
    out: list[str] = []
    for page in text.replace("\r\n", "\n").replace("\r", "\n").split(PAGE_BREAK):
        lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in page.split("\n")]
        merged: list[str] = []
        for ln in lines:
            if (
                merged
                and ln
                and len(merged[-1]) >= 45  # wrapped lines run to the edge of the column; short lines are headers or contact details
                and not _HEADING_RE.match(ln)
                and not _HEADING_RE.match(merged[-1])
                and not _BULLET_RE.match(ln)
                and ln[0].islower()
                and merged[-1][-1] not in ".:;!?"
            ):
                merged[-1] += " " + ln
            else:
                merged.append(ln)
        spaced: list[str] = []
        for ln in merged:
            if _HEADING_RE.match(ln) and spaced and spaced[-1]:
                spaced.append("")
            spaced.append(ln)
        out.append(re.sub(r"\n{3,}", "\n\n", "\n".join(spaced)).strip("\n"))
    return ("\n" + PAGE_BREAK + "\n").join(out)


def score_text(text: str, pages: int) -> tuple[int, list[str]]:
    """0-100 parse quality, plus plain-language warnings."""
    warnings: list[str] = []
    body = text.replace(PAGE_BREAK, "\n")
    chars = len(body.strip())
    if chars < 40 * max(pages, 1):
        warnings.append("Almost no text was found. This looks like a scanned or image-only file; upload a text-based PDF or DOCX.")
        return min(10, chars // 5), warnings
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", body)
    letters = sum(c.isalpha() for c in body)
    odd = sum(1 for c in body if c == "�" or (ord(c) < 32 and c not in "\n\t"))
    score = 100.0
    if chars < 600:
        score -= 25
        warnings.append("Very little text was extracted; check that nothing is missing.")
    if letters / max(chars, 1) < 0.55:
        score -= 25
        warnings.append("Much of the text is symbols or digits, which can mean the layout was garbled.")
    if odd / max(chars, 1) > 0.01:
        score -= 25
        warnings.append("Unreadable characters were found.")
    if words:
        avg = sum(map(len, words)) / len(words)
        single = sum(1 for w in words if len(w) == 1 and w.lower() not in "ai") / len(words)
        if avg < 3.0 or single > 0.25:
            score -= 30
            warnings.append("Words look split into single letters, so columns or fonts may have been read badly.")
    lines = [ln for ln in body.split("\n") if ln.strip()]
    if len(lines) < 8 and chars > 1500:
        score -= 15
        warnings.append("Line breaks were lost; the structure may be hard to recover.")
    return max(0, min(100, round(score))), warnings


def parse_bytes(data: bytes, name: str) -> Parsed:
    kind = detect_kind(data, name)
    pages = 1
    warnings: list[str] = []
    if kind == "pdf":
        text, pages, parser = "", 1, "pymupdf"
        try:
            text, pages = _pdf_pymupdf(data)
        except HTTPException:
            raise
        except ImportError:
            parser = "none"
        except Exception:
            parser = "none"
        text = clean(text)
        q, warnings = score_text(text, pages)
        if q < 40 or parser == "none":
            try:
                alt, alt_pages = _pdf_pdfplumber(data)
                alt = clean(alt)
                aq, aw = score_text(alt, alt_pages)
                if aq > q or parser == "none":
                    text, pages, q, warnings = alt, alt_pages, aq, aw
            except Exception:
                if parser == "none":
                    raise HTTPException(400, "That PDF could not be read.") from None
        return Parsed(text, q, warnings, pages, kind)
    if kind == "docx":
        text = clean(_docx(data))
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        text = clean(text)
    q, warnings = score_text(text, pages)
    return Parsed(text, q, warnings, pages, kind)


def page_of_line(text: str, line_no: int) -> int | None:
    """1-based page of a 1-based line number, or None when the text has no page breaks."""
    if PAGE_BREAK not in text:
        return None
    lines = text.split("\n")
    return 1 + sum(ln.count(PAGE_BREAK) for ln in lines[: max(0, line_no - 1)])


def numbered_lines(text: str) -> list[tuple[int, str]]:
    """(1-based line number, text) for non-empty lines; the page break character is stripped."""
    return [(i + 1, ln.replace(PAGE_BREAK, "").strip()) for i, ln in enumerate(text.split("\n")) if ln.replace(PAGE_BREAK, "").strip()]
